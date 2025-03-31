# import openpyxl
#
# # def get_data():
# # 定义excel文件路径
# excel_file=r"F:\testui\datas\baidu.xlsx"
# # 创建读取对象
# wb=openpyxl.load_workbook(excel_file)
# # 获取所有工作表
# all_sheet=wb.worksheets
# print(all_sheet)
# # ws=all_sheet
# # 指定工作表为第1张表
# s1=all_sheet[0]
# # return s1
# # 获取第1张表的第1行
# print("表格第1行：",s1[1])
# for i in s1[1]:
#     print(i.value)

from openpyxl import load_workbook

# def read_excel(file_path):
#     wb = load_workbook(file_path, read_only=True)
#     sheet = wb.active
#
#     # 读取标题
#     headers = [cell.value for cell in sheet[1]]
#
#     # 读取数据
#     data = []
#     for row in sheet.iter_rows(min_row=2, values_only=True):
#         data.append(dict(zip(row[0],(row[1],row[2]))))
#
#     wb.close()
#     return data

from openpyxl import load_workbook


def excel_to_locator_dict(file_path, sheet_name=None):
    """
    从 Excel 中读取数据并构建定位字典
    :param file_path: Excel 文件路径
    :param sheet_name: 工作表名称（默认第一个工作表）
    :return: {名称: (定位方式, 元素)} 的字典
    """
    wb = load_workbook(file_path, read_only=True)
    sheet = wb[sheet_name] if sheet_name else wb.active

    # 获取标题列的索引（假设第一行是标题）
    headers = [cell.value for cell in sheet[1]]  # 读取第一行
    column_map = {name: idx + 1 for idx, name in enumerate(headers)}  # 列号从1开始

    # 校验必要列是否存在
    required_columns = ["名称", "定位方式", "元素"]
    for col in required_columns:
        if col not in headers:
            wb.close()
            raise ValueError(f"缺少必要列: {col}")

    # 构建字典
    locator_dict = {}
    for row in sheet.iter_rows(min_row=2):  # 从第二行开始遍历
        name = row[column_map["名称"] - 1].value  # 转换为0-based索引
        by_type = row[column_map["定位方式"] - 1].value
        element = row[column_map["元素"] - 1].value

        if not all([name, by_type, element]):
            continue  # 跳过空行

        locator_dict[name] = (by_type, element)

    wb.close()
    return locator_dict


# 使用示例 ----------------------
if __name__ == "__main__":
    locators = excel_to_locator_dict(r"F:\testbaiduui\datas\baidu.xlsx")
    print(locators)
    # 查询示例
    print(locators.get("butten_baiduyixia", "未找到"))  # 输出: ('By.ID', 'su')
    print(locators.get("PYTHON_PATH", "未找到"))  # 输出: ('By.LINK_TEXT', 'Welcome to Python.org')